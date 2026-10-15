"""Populate the single SmartLab workbook used by Spec Verification Lab.

The workbook remains in the teacher-provided 10-sheet format.  This script
changes only specification values; it deliberately preserves the workbook's
existing sheet names, order, formatting, and dimensions where possible.
"""

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "spec" / "smartlab_access_controller_spec.xlsx"


def replace_rows(ws, rows):
    """Replace data rows while retaining the template header row and its style."""
    for row in range(2, ws.max_row + 1):
        for column in range(1, ws.max_column + 1):
            ws.cell(row, column).value = None
    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            ws.cell(row_index, column_index).value = value
    final_row = len(rows) + 1
    if ws.max_row > final_row:
        ws.delete_rows(final_row + 1, ws.max_row - final_row)


def main():
    workbook = load_workbook(WORKBOOK)

    workbook["ThongTin"]["B3"] = (
        "Bộ điều khiển mở cửa SmartLab xác thực thẻ và mã mở khóa khẩn cấp. "
        "Access worker và audit worker dùng chung access_guard; hệ thống ghi "
        "nhận sự kiện truy cập và dùng các fixture lỗi có chủ đích cho thực hành."
    )

    replace_rows(
        workbook["YeuCauChucNang"],
        [
            ("REQ-1", "Access worker và audit worker không được đồng thời giữ access_guard.", "safety", "HIGH"),
            ("REQ-2", "Cụm access_guard không được deadlock khi hai worker tranh chấp khóa.", "liveness", "HIGH"),
            ("REQ-3", "Một phiên xét duyệt thông thường phải hoàn tất sau khi cấp quyền hoặc từ chối.", "liveness", "HIGH"),
            ("REQ-4", "Hàm phân tích mã thẻ không được sao chép quá bộ đệm 20 ký tự ASCII.", "memory", "HIGH"),
            ("REQ-5", "Hàm tạo audit record phải giải phóng vùng nhớ đã cấp phát.", "memory", "MEDIUM"),
            ("REQ-6", "Hàm đóng phiên không được giải phóng cùng một con trỏ nhiều hơn một lần.", "memory", "HIGH"),
            ("REQ-7", "Hàm xóa badge cache không được truy cập vùng nhớ sau khi giải phóng.", "memory", "HIGH"),
            ("REQ-8", "Bộ đếm sự kiện truy cập không được tràn int32 khi cộng các bộ đếm đầu vào.", "safety", "MEDIUM"),
            ("REQ-9", "Hàm tỷ lệ truy cập hợp lệ phải xử lý mẫu số bằng 0 có kiểm soát.", "functional", "MEDIUM"),
            ("REQ-10", "Mã mở cửa khẩn cấp gồm hai số nguyên 16-bit phải được kiểm thử bằng black-box và concolic fuzzing.", "functional", "HIGH"),
        ],
    )

    replace_rows(
        workbook["MoHinhTrangThai"],
        [
            ("access_free", "x", "safe", "access_holds, audit_holds, entry_idle"),
            ("access_holds", None, "access_locked", "access_conflict"),
            ("audit_holds", None, "audit_locked", "access_conflict"),
            ("access_conflict", None, "access_locked, audit_locked, violation", "entry_idle"),
            ("entry_idle", None, "entry_ready, door_locked", "credential_review, override_review"),
            ("credential_review", None, "normal_entry_active", "access_granted, access_denied"),
            ("override_review", None, "override_entry_active", "access_granted, access_denied"),
            ("access_granted", None, "door_unlocked", "session_closed"),
            ("access_denied", None, "door_locked", "session_closed"),
            ("session_closed", None, "complete", "session_closed"),
        ],
    )

    replace_rows(
        workbook["ThuocTinhAnToan"],
        [
            ("AG(!violation)",),
            ("EF(access_locked)",),
            ("AF(complete)",),
        ],
    )

    replace_rows(
        workbook["LoaiTruLanNhau"],
        [
            ("access_locked", "audit_locked"),
            ("violation", "safe"),
            ("access_locked", "safe"),
            ("audit_locked", "safe"),
            ("normal_entry_active", "override_entry_active"),
        ],
    )

    replace_rows(
        workbook["HamSinhMa"],
        [
            ("parse_badge_code", "c", "buffer_copy", "raw_badge", 20, "buffer-overflow", None, None),
            ("create_audit_record", "c", "alloc_release", None, 32, "memory-leak", None, None),
            ("close_access_session", "c", "alloc_release", None, 8, "double-free", None, None),
            ("reset_badge_cache", "c", "alloc_release", None, 16, "use-after-free", None, None),
            ("access_event_total", "c", "arithmetic", "granted_events, pending_events", 16, None, None, None),
            ("valid_access_ratio", "python", "divide", "granted, attempted", 16, None, None, None),
            ("unlock_override_code", "python", "magic_value", "override_a, override_b", 16, None, 5813, 2649),
        ],
    )

    replace_rows(
        workbook["KiemTraSoHoc"],
        [
            (
                "access_event_total",
                "granted_events, pending_events",
                "granted_events >= 0 and pending_events >= 0 and granted_events + pending_events > 2000000000",
                "granted_events + pending_events",
            ),
            (
                "audit_byte_total",
                "written_bytes, buffered_bytes",
                "written_bytes >= 0 and buffered_bytes >= 0 and written_bytes + buffered_bytes > 2000000000",
                "written_bytes + buffered_bytes",
            ),
            (
                "daily_badge_total",
                "accepted_today, queued_today",
                "accepted_today >= 0 and queued_today >= 0 and accepted_today + queued_today > 32767",
                "accepted_today + queued_today",
            ),
            (
                "active_session_total",
                "active_sessions, closing_sessions",
                "active_sessions >= 0 and closing_sessions >= 0 and active_sessions + closing_sessions > 65535",
                "active_sessions + closing_sessions",
            ),
            (
                "denied_attempt_total",
                "denied_today, delayed_denials",
                "denied_today >= 0 and delayed_denials >= 0 and denied_today + delayed_denials > 32767",
                "denied_today + delayed_denials",
            ),
        ],
    )

    replace_rows(
        workbook["TestCase"],
        [
            ("TC-1", "REQ-1, REQ-2", "Cụm khóa access_guard cố ý có xung đột", "Hệ thống ở access_free", "Chạy kiểm chứng mô hình\nKiểm tra trạng thái access_conflict", "AG(!violation) vi phạm và mutex phát hiện trạng thái nguy hiểm có chủ đích", "HIGH"),
            ("TC-2", "REQ-3", "Vòng đời một phiên xét duyệt", "Hệ thống ở access_free", "Chạy kiểm chứng mô hình\nKiểm tra CTL liveness", "AF(complete) đúng từ trạng thái khởi tạo", "HIGH"),
            ("TC-3", "REQ-4, CON-2", "Mã thẻ dài hơn 20 ký tự", "Hàm parse_badge_code được sinh với test fixture", "Sinh mã C\nChạy static checker", "Phát hiện dangerous-function (CWE-120) trên test fixture", "HIGH"),
            ("TC-4", "REQ-5, REQ-6, REQ-7", "Vòng đời bộ nhớ audit và phiên", "Ba hàm alloc_release được sinh với fixture", "Sinh mã C\nChạy static checker", "Phát hiện CWE-401, CWE-415 và CWE-416 trên fixture", "HIGH"),
            ("TC-5", "REQ-8", "Bộ đếm sự kiện vượt ngưỡng int32", "Hai biến đếm không âm", "Chạy kiểm tra SMT/Z3", "SAT: tồn tại đầu vào có thể tràn số", "MEDIUM"),
            ("TC-6", "REQ-9", "Mẫu số tỷ lệ truy cập bằng 0", "Hàm valid_access_ratio được sinh", "Chạy black-box fuzzing 300 lần", "Có ít nhất một ZeroDivisionError trên fixture", "MEDIUM"),
            ("TC-7", "REQ-10, CON-3", "Tìm mã mở cửa khẩn cấp", "Hàm unlock_override_code có cặp test 5813, 2649", "Chạy concolic fuzzing tối đa 15 vòng", "Tìm đúng cặp và đạt nhánh test oracle", "HIGH"),
            ("TC-8", "REQ-10, NFR-1, NFR-2, NFR-3", "So sánh fuzzing trên target của nhóm", "Đã chốt seed, ngân sách và môi trường", "Chạy 5 seed black-box\nChạy concolic cùng nhánh SLAC\nLưu CSV và JSON", "Bảng có mean, median, thời gian và input đạt đích", "HIGH"),
        ],
    )

    workbook.save(WORKBOOK)


if __name__ == "__main__":
    main()
