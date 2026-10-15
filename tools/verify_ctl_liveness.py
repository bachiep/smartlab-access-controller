"""Verify the full liveness formula that the teacher tool cannot parse."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from slac_fuzz.ctl import check_ag_implies_af


def main() -> None:
    workbook = load_workbook(ROOT / "spec" / "smartlab_access_controller_spec.xlsx", data_only=True)
    rows = list(workbook["MoHinhTrangThai"].iter_rows(min_row=2, values_only=True))
    initial = {row[0] for row in rows if row[1] == "x"}
    transitions = {row[0]: {state.strip() for state in row[3].split(",")} for row in rows}
    labels = {row[0]: {label.strip() for label in row[2].split(",")} for row in rows}

    result = check_ag_implies_af(
        initial, transitions, labels, "normal_entry_active", "complete"
    )
    payload = {
        "formula": "AG(normal_entry_active -> AF(complete))",
        "initial_states": sorted(initial),
        "holds": result.holds,
        "violating_states": sorted(result.violating_states),
        "satisfying_states": sorted(result.satisfying_states),
        "note": "Independent project checker; Spec Verification Lab evaluates only simple atomic CTL formulas.",
    }
    output = ROOT / "results" / "ctl_liveness_check.json"
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    if not result.holds:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
