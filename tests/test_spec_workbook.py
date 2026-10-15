from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "spec" / "smartlab_access_controller_spec.xlsx"
EXPECTED_SHEETS = [
    "ThongTin",
    "YeuCauChucNang",
    "YeuCauPhiChucNang",
    "RangBuoc",
    "MoHinhTrangThai",
    "ThuocTinhAnToan",
    "LoaiTruLanNhau",
    "HamSinhMa",
    "KiemTraSoHoc",
    "TestCase",
]


def test_workbook_keeps_the_teacher_template_sheet_order():
    workbook = load_workbook(WORKBOOK, data_only=False)
    assert workbook.sheetnames == EXPECTED_SHEETS


def test_workbook_contains_the_required_educational_fault_fixtures():
    workbook = load_workbook(WORKBOOK, data_only=False)
    code_rows = list(workbook["HamSinhMa"].iter_rows(min_row=2, values_only=True))
    faults = {row[5] for row in code_rows if row[5]}

    assert faults == {
        "buffer-overflow",
        "memory-leak",
        "double-free",
        "use-after-free",
    }
    assert {row[2] for row in code_rows} >= {
        "buffer_copy",
        "alloc_release",
        "divide",
        "magic_value",
    }


def test_workbook_models_intentional_conflict_and_completed_denial_path():
    workbook = load_workbook(WORKBOOK, data_only=False)
    transitions = {
        row[0]: row[3]
        for row in workbook["MoHinhTrangThai"].iter_rows(min_row=2, values_only=True)
    }
    labels = {
        row[0]: row[2]
        for row in workbook["MoHinhTrangThai"].iter_rows(min_row=2, values_only=True)
    }

    assert "violation" in labels["access_conflict"]
    assert transitions["access_denied"] == "session_closed"


def test_workbook_has_one_initial_state_and_a_terminal_completion_state():
    workbook = load_workbook(WORKBOOK, data_only=False)
    rows = list(workbook["MoHinhTrangThai"].iter_rows(min_row=2, values_only=True))

    assert [row[0] for row in rows if row[1] == "x"] == ["access_free"]
    transitions = {row[0]: row[3] for row in rows}
    assert transitions["session_closed"] == "session_closed"


def test_liveness_uses_the_tool_supported_af_formula():
    workbook = load_workbook(WORKBOOK, data_only=False)
    formulas = [row[0] for row in workbook["ThuocTinhAnToan"].iter_rows(min_row=2, values_only=True)]

    assert "AF(complete)" in formulas


def test_every_nonterminal_state_moves_toward_completion():
    workbook = load_workbook(WORKBOOK, data_only=False)
    transitions = {
        row[0]: [name.strip() for name in row[3].split(",")]
        for row in workbook["MoHinhTrangThai"].iter_rows(min_row=2, values_only=True)
    }

    nonterminal = set(transitions) - {"session_closed"}
    remaining = set(nonterminal)
    while remaining:
        removable = {
            state
            for state in remaining
            if all(next_state == "session_closed" or next_state not in remaining for next_state in transitions[state])
        }
        assert removable, "a nonterminal transition cycle can avoid completion"
        remaining -= removable
