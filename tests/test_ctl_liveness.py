from pathlib import Path

from openpyxl import load_workbook

from slac_fuzz.ctl import check_ag_implies_af


WORKBOOK = Path(__file__).resolve().parents[1] / "spec" / "smartlab_access_controller_spec.xlsx"


def _model_from_workbook():
    workbook = load_workbook(WORKBOOK, data_only=True)
    rows = list(workbook["MoHinhTrangThai"].iter_rows(min_row=2, values_only=True))
    return (
        {row[0] for row in rows if row[1] == "x"},
        {row[0]: {target.strip() for target in row[3].split(",")} for row in rows},
        {row[0]: {label.strip() for label in row[2].split(",")} for row in rows},
    )


def test_nested_liveness_formula_holds_for_the_group_model():
    initial, transitions, labels = _model_from_workbook()

    result = check_ag_implies_af(
        initial, transitions, labels, "normal_entry_active", "complete"
    )

    assert result.holds is True
    assert result.violating_states == frozenset()


def test_nested_liveness_reports_a_path_that_can_avoid_completion():
    result = check_ag_implies_af(
        {"review"},
        {"review": {"review", "complete"}, "complete": {"complete"}},
        {"review": {"normal_entry_active"}, "complete": {"complete"}},
        "normal_entry_active",
        "complete",
    )

    assert result.holds is False
    assert result.violating_states == frozenset({"review"})
